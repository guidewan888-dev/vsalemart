"""Build an idempotent Supabase catalog migration from Shopee Seller Centre exports.

Expected exports: basic info, sales info, shipping info, days-to-ship, and media.
Shopee puts machine-readable headers in row 1 and product rows from row 7 onward.
"""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl.worksheet.views import Pane

# Shopee currently emits bottom_left, while openpyxl only whitelists bottomLeft.
Pane.activePane.values.add("bottom_left")

from openpyxl import load_workbook  # noqa: E402

SHOP_ID = "501468956"
IMPORT_KEY = "shopee-501468956-2026-09-07"


def sql_text(value: Any) -> str:
    if value is None:
        return "null"
    normalized = str(value).replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    normalized = "\n".join(line.rstrip() for line in normalized.split("\n")).strip()
    return "'" + normalized.replace("'", "''") + "'"


def sql_number(value: Any, default: str = "0") -> str:
    if value in (None, ""):
        return default
    try:
        return str(Decimal(str(value)))
    except InvalidOperation:
        return default


def as_int(value: Any) -> int:
    try:
        return max(0, int(Decimal(str(value or 0))))
    except (InvalidOperation, ValueError):
        return 0


def workbook_rows(path: Path) -> list[dict[str, Any]]:
    sheet = load_workbook(path, data_only=True, read_only=False).active
    headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
    rows: list[dict[str, Any]] = []
    for row_number in range(7, sheet.max_row + 1):
        row = {str(header): sheet.cell(row_number, index + 1).value for index, header in enumerate(headers) if header}
        if row.get("et_title_product_id"):
            rows.append(row)
    return rows


def find_export(source: Path, token: str) -> Path:
    matches = sorted(source.glob(f"*_{token}_*.xlsx"))
    if not matches:
        raise FileNotFoundError(f"Missing Shopee {token} export in {source}")
    return matches[-1]


def category_slug(name: str, path: str) -> str:
    combined = f"{name} {path}".lower()
    if any(word in combined for word in ("ธง", "เสาธง", "flag")):
        return "flags-ceremony"
    if any(word in combined for word in ("ปพ.", "แบบบันทึก", "แบบประเมิน", "ทะเบียน", "ระเบียน", "ประกาศนียบัตร")):
        return "forms-documents"
    if any(word in combined for word in ("sport", "กีฬา", "ฟุตบอล", "ลูกบอล")):
        return "sports"
    if any(word in combined for word in ("music", "ดนตรี", "เมโลเดียน", "ขลุ่ย")):
        return "music"
    if "books & magazines" in combined or any(word in combined for word in ("หนังสือเรียน", "แบบเรียน", "workbook")):
        return "books-workbooks"
    if any(word in combined for word in ("art supplies", "craft", "ศิลปะ", "วาดเขียน", "สีน้ำ", "สีโปสเตอร์")):
        return "art-craft"
    if any(word in combined for word in ("adhesives", "tape", "กาว", "เทป")):
        return "tape-adhesive"
    if any(word in combined for word in ("pen", "pencil", "eraser", "marker", "ปากกา", "ดินสอ", "ยางลบ", "เครื่องเขียน")):
        return "stationery"
    return "office-supplies"


def chunks(items: list[str], size: int = 150) -> list[list[str]]:
    return [items[index:index + size] for index in range(0, len(items), size)]


def insert_values(table: str, columns: str, values: list[str], conflict: str) -> str:
    statements = []
    for batch in chunks(values):
        statements.append(f"insert into {table} ({columns}) values\n" + ",\n".join(batch) + f"\n{conflict};")
    return "\n\n".join(statements)


def build(source: Path, output: Path) -> dict[str, int]:
    basic_rows = workbook_rows(find_export(source, "basic_info"))
    dts_rows = workbook_rows(find_export(source, "dts_info"))
    media_rows = workbook_rows(find_export(source, "media_info"))
    sales_rows = workbook_rows(find_export(source, "sales_info"))
    shipping_rows = workbook_rows(find_export(source, "shipping_info"))

    basic = {str(row["et_title_product_id"]): row for row in basic_rows}
    media = {str(row["et_title_product_id"]): row for row in media_rows}
    shipping = {
        (str(row["et_title_product_id"]), str(row["et_title_variation_id"])): row
        for row in shipping_rows
    }
    sales_by_product: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in sales_rows:
        sales_by_product[str(row["et_title_product_id"])].append(row)

    product_values: list[str] = []
    image_values: list[str] = []
    variant_values: list[str] = []
    active_count = 0

    for product_index, row in enumerate(dts_rows):
        product_id = str(row["et_title_product_id"])
        detail = basic.get(product_id, {})
        media_row = media.get(product_id, {})
        variants = sales_by_product.get(product_id, [])
        prices = [Decimal(str(item["et_title_variation_price"])) for item in variants if item.get("et_title_variation_price") not in (None, "")]
        stock = sum(as_int(item.get("et_title_variation_stock")) for item in variants)
        name = str(row.get("et_title_product_name") or detail.get("et_title_product_name") or f"สินค้า {product_id}").strip()
        description = str(detail.get("et_title_product_description") or "").strip()
        category_path = str(row.get("et_title_product_category") or media_row.get("et_title_product_category") or "").strip()
        parent_sku = row.get("et_title_parent_sku") or detail.get("et_title_parent_sku")
        cover = media_row.get("ps_item_cover_image")
        is_active = bool(stock > 0 and prices and cover)
        active_count += int(is_active)
        price = min(prices) if prices else Decimal("0")
        product_values.append(
            "(" + ",".join([
                sql_text(f"shopee-{product_id}"), sql_text(name), sql_text(description), sql_number(price),
                "(select id from public.categories where slug=" + sql_text(category_slug(name, category_path)) + ")",
                sql_text(cover), sql_number(stock), sql_text("ready" if is_active else None),
                "true" if product_index < 60 and is_active else "false", "false", "false",
                "true" if is_active else "false", sql_text("shopee"), sql_text(product_id),
                sql_text(f"https://shopee.co.th/product/{SHOP_ID}/{product_id}"), sql_text(parent_sku),
                sql_text(category_path), sql_number(row.get("et_title_product_dts"), "null"), "now()"
            ]) + ")"
        )

        seen_images: set[str] = set()
        image_urls = [media_row.get("ps_item_cover_image")] + [media_row.get(f"ps_item_image.{index}") for index in range(1, 9)]
        image_urls += [media_row.get(key) for key in media_row if key.startswith("et_title_option_image_")]
        for sort_order, url in enumerate(url for url in image_urls if url and url not in seen_images):
            seen_images.add(str(url))
            image_values.append(
                "((select id from public.products where source='shopee' and source_product_id="
                + sql_text(product_id) + ")," + sql_text(url) + "," + sql_text(name) + f",{sort_order})"
            )

        for variant in variants:
            variant_id = str(variant.get("et_title_variation_id") or "0")
            ship = shipping.get((product_id, variant_id), {})
            shipping_options = {
                key.removeprefix("et_title_product_channel_toggle."): value
                for key, value in ship.items()
                if key.startswith("et_title_product_channel_toggle.") and value is not None
            }
            variant_values.append(
                "(" + ",".join([
                    "(select id from public.products where source='shopee' and source_product_id=" + sql_text(product_id) + ")",
                    sql_text(variant_id), sql_text(variant.get("et_title_variation_name")),
                    sql_text(variant.get("et_title_variation_sku")), sql_number(variant.get("et_title_variation_price")),
                    sql_number(as_int(variant.get("et_title_variation_stock"))), sql_text(variant.get("ps_gtin_code")),
                    sql_number(ship.get("et_title_product_weight"), "null"), sql_number(ship.get("et_title_product_length"), "null"),
                    sql_number(ship.get("et_title_product_width"), "null"), sql_number(ship.get("et_title_product_height"), "null"),
                    sql_number(variant.get("ps_minimum_purchase_quantity"), "null"), sql_number(variant.get("ps_maximum_purchase_quantity"), "null"),
                    sql_text(json.dumps(shipping_options, ensure_ascii=False)) + "::jsonb"
                ]) + ")"
            )

    sql = [
        "-- Generated from the five official Shopee Seller Centre exports.",
        "begin;",
        "update public.products set is_active=false, updated_at=now() where source='shopee';",
        insert_values(
            "public.products",
            "slug,name,description,price,category_id,cover_image,stock,badge,is_featured,is_flash_sale,is_demo,is_active,source,source_product_id,source_url,parent_sku,category_path,preparation_days,synced_at",
            product_values,
            "on conflict (source,source_product_id) do update set slug=excluded.slug,name=excluded.name,description=excluded.description,price=excluded.price,category_id=excluded.category_id,cover_image=excluded.cover_image,stock=excluded.stock,badge=excluded.badge,is_featured=excluded.is_featured,is_flash_sale=excluded.is_flash_sale,is_demo=false,is_active=excluded.is_active,source_url=excluded.source_url,parent_sku=excluded.parent_sku,category_path=excluded.category_path,preparation_days=excluded.preparation_days,synced_at=excluded.synced_at,updated_at=now()"
        ),
        "delete from public.product_images where product_id in (select id from public.products where source='shopee');",
        insert_values(
            "public.product_images", "product_id,url,alt_text,sort_order", image_values,
            "on conflict (product_id,url) do update set alt_text=excluded.alt_text,sort_order=excluded.sort_order"
        ),
        "delete from public.product_variants where product_id in (select id from public.products where source='shopee');",
        insert_values(
            "public.product_variants",
            "product_id,source_variant_id,name,sku,price,stock,gtin,weight_kg,length_cm,width_cm,height_cm,minimum_purchase_quantity,maximum_purchase_quantity,shipping_options",
            variant_values,
            "on conflict (product_id,source_variant_id) do update set name=excluded.name,sku=excluded.sku,price=excluded.price,stock=excluded.stock,gtin=excluded.gtin,weight_kg=excluded.weight_kg,length_cm=excluded.length_cm,width_cm=excluded.width_cm,height_cm=excluded.height_cm,minimum_purchase_quantity=excluded.minimum_purchase_quantity,maximum_purchase_quantity=excluded.maximum_purchase_quantity,shipping_options=excluded.shipping_options,updated_at=now()"
        ),
        "update public.products set is_active=false where is_demo;",
        "insert into public.catalog_sync_runs (source,source_shop_id,import_key,status,product_count,active_product_count,variant_count,image_count,metadata,completed_at) values ("
        + ",".join([
            sql_text("shopee"), sql_text(SHOP_ID), sql_text(IMPORT_KEY), sql_text("completed"), str(len(product_values)),
            str(active_count), str(len(variant_values)), str(len(image_values)),
            sql_text(json.dumps({"exported_at": "2026-09-07", "method": "Seller Centre mass update exports"})) + "::jsonb", "now()"
        ]) + ") on conflict (import_key) do update set status=excluded.status,product_count=excluded.product_count,active_product_count=excluded.active_product_count,variant_count=excluded.variant_count,image_count=excluded.image_count,metadata=excluded.metadata,completed_at=excluded.completed_at;",
        "commit;",
    ]
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("\n\n".join(sql) + "\n", encoding="utf-8")
    return {"products": len(product_values), "active_products": active_count, "variants": len(variant_values), "images": len(image_values)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=Path("data/shopee-export/2026-09-07"))
    parser.add_argument("--output", type=Path, default=Path("supabase/migrations/20260907021000_import_shopee_catalog.sql"))
    args = parser.parse_args()
    stats = build(args.source, args.output)
    print(json.dumps(stats, ensure_ascii=False))
